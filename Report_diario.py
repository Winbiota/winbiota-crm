import os, json, datetime, smtplib
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build

creds_json = json.loads(os.environ['GOOGLE_CREDENTIALS'])
SHEET_ID   = os.environ['SHEET_ID']
EMAIL_USER = os.environ['EMAIL_USER']
EMAIL_PASS = os.environ['EMAIL_PASS']
OUTPUT     = '/tmp/Winbiota_Reporte_Diario.xlsx'

scopes = ['https://www.googleapis.com/auth/spreadsheets']
creds  = Credentials.from_service_account_info(creds_json, scopes=scopes)
sheets = build('sheets', 'v4', credentials=creds)

# ── Read snapshots ────────────────────────────────────────────────
snap_result = sheets.spreadsheets().values().get(
    spreadsheetId=SHEET_ID,
    range='Snapshots!A2:G5000'
).execute()
snap_rows = snap_result.get('values', [])
print(f"Snapshots: {len(snap_rows)} filas")

# ── Read CRM for estatus and leads efectivos ──────────────────────
crm_result = sheets.spreadsheets().values().get(
    spreadsheetId=SHEET_ID,
    range='CRM Winbiota',
    valueRenderOption='FORMATTED_VALUE',
    dateTimeRenderOption='FORMATTED_STRING'
).execute()
crm_rows = crm_result.get('values', [])
max_cols = max(len(r) for r in crm_rows)
crm_rows = [r + [''] * (max_cols - len(r)) for r in crm_rows]

buenos_vals1 = ['CONTESTO','CONTESTA','INTERES','LLAMADA','NO INT','SI']
buenos_vals2 = ['SI CONTESTA','INTERESADA']

# Process CRM data
estatus1_counts = {}
estatus2_counts = {}
leads_efectivos = []

def parse_date(val):
    v = str(val).strip()
    for fmt in ('%d/%m/%Y','%d/%m/%y','%Y-%m-%d'):
        try: return datetime.datetime.strptime(v, fmt).strftime('%d/%m/%Y')
        except: pass
    return v

for row in crm_rows[3:]:
    name = str(row[2]).strip() if len(row) > 2 else ''
    if 'dummy data' in name.lower() or name == '': continue
    fecha = parse_date(row[1]) if len(row) > 1 else ''
    tel   = str(row[4]).replace('p:','').strip() if len(row) > 4 else ''
    e1    = str(row[12]).strip() if len(row) > 12 else ''
    e2    = str(row[16]).strip() if len(row) > 16 else ''
    c1    = str(row[11]).strip() if len(row) > 11 else ''
    c2    = str(row[15]).strip() if len(row) > 15 else ''
    if c1 and c1 not in ['','nan','None']:
        k = e1 if e1 and e1 not in ['','nan','None'] else 'SIN ESTATUS'
        estatus1_counts[k] = estatus1_counts.get(k, 0) + 1
        if e1.upper() in buenos_vals1:
            leads_efectivos.append([fecha, name, tel, e1, 'LL1'])
    if c2 and c2 not in ['','nan','None']:
        k = e2 if e2 and e2 not in ['','nan','None'] else 'SIN ESTATUS'
        estatus2_counts[k] = estatus2_counts.get(k, 0) + 1
        if e2.upper() in buenos_vals2:
            leads_efectivos.append([fecha, name, tel, e2, 'LL2'])

# ── Build Excel ───────────────────────────────────────────────────
G_D='1B5E20'; G_M='2E7D32'; G_L='E8F5E9'
WHITE='FFFFFF'; BLUE_D='1565C0'; AMBER='FFF8E1'; GRAY='F5F5F5'

def st(cell, bold=False, sz=9, col="000000", bg=None, ha='left', va='center', wrap=False):
    cell.font = Font(name='Arial', bold=bold, size=sz, color=col)
    if bg: cell.fill = PatternFill('solid', start_color=bg)
    cell.alignment = Alignment(horizontal=ha, vertical=va, wrap_text=wrap)

thin = Side(style='thin', color='C8E6C9')
bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
thin2 = Side(style='thin', color='B0BEC5')
bdr2 = Border(left=thin2, right=thin2, top=thin2, bottom=thin2)

def hrow(ws, row, labels, widths, bg=G_M):
    for i,(h,w) in enumerate(zip(labels,widths),1):
        c = ws.cell(row=row, column=i, value=h)
        st(c, bold=True, sz=9, col=WHITE, bg=bg, ha='center', wrap=True)
        c.border = bdr
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 28

def trow(ws, row, text, ncols, bg=G_D):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    st(c, bold=True, sz=12, col=WHITE, bg=bg, ha='center')
    ws.row_dimensions[row].height = 26

def section(ws, row, text, ncols, bg=BLUE_D):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    st(c, bold=True, sz=10, col=WHITE, bg=bg, ha='left')
    ws.row_dimensions[row].height = 20

wb = openpyxl.Workbook()

# ── HOJA 1: FRANJAS HORARIAS POR DIA ─────────────────────────────
# Group snapshots by date
from collections import defaultdict
by_date = defaultdict(list)
for row in snap_rows:
    if len(row) >= 7:
        date = row[1]
        hour = row[2]
        try:
            ll1 = int(row[3]); ll1b = int(row[4])
            ll2 = int(row[5]); ll2b = int(row[6])
            by_date[date].append((hour, ll1, ll1b, ll2, ll2b))
        except: pass

ws1 = wb.active
ws1.title = "Franjas Horarias"
NCOLS = 9
today = datetime.datetime.utcnow().strftime('%d/%m/%Y')
trow(ws1, 1, f'WINBIOTA - Llamadas por Franja Horaria', NCOLS)
r = 3

for date in sorted(by_date.keys()):
    snaps = sorted(by_date[date], key=lambda x: x[0])
    section(ws1, r, f'Dia: {date}', NCOLS); r+=1
    hrow(ws1, r,
         ['Franja','LL1 nuevas','LL1 efectivas','% Efect LL1','LL2 nuevas','LL2 efectivas','% Efect LL2','LL tot nuevas','LL tot efect'],
         [14,12,14,12,12,14,12,14,14]); r+=1
    prev_ll1=prev_ll1b=prev_ll2=prev_ll2b=0
    day_ll1=day_ll1b=day_ll2=day_ll2b=0
    for i,(hour, ll1, ll1b, ll2, ll2b) in enumerate(snaps):
        new_ll1  = ll1  - prev_ll1
        new_ll1b = ll1b - prev_ll1b
        new_ll2  = ll2  - prev_ll2
        new_ll2b = ll2b - prev_ll2b
        prev_ll1=ll1; prev_ll1b=ll1b; prev_ll2=ll2; prev_ll2b=ll2b
        day_ll1+=new_ll1; day_ll1b+=new_ll1b; day_ll2+=new_ll2; day_ll2b+=new_ll2b
        # Franja label
        try:
            h = int(hour.split(':')[0])
            franja = f'{h:02d}:00 - {h+1:02d}:00'
        except: franja = hour
        pct1 = f"{round(new_ll1b/new_ll1*100,1)}%" if new_ll1>0 else '0%'
        pct2 = f"{round(new_ll2b/new_ll2*100,1)}%" if new_ll2>0 else '0%'
        tot_new = new_ll1+new_ll2; tot_b = new_ll1b+new_ll2b
        bg = G_L if r%2==0 else WHITE
        vals = [franja, new_ll1, new_ll1b, pct1, new_ll2, new_ll2b, pct2, tot_new, tot_b]
        for ci,val in enumerate(vals,1):
            c = ws1.cell(row=r, column=ci, value=val)
            st(c, sz=9, bg=bg, ha='center' if ci>1 else 'left'); c.border = bdr2
        ws1.row_dimensions[r].height = 16; r+=1
    # Day total
    pct1t = f"{round(day_ll1b/day_ll1*100,1)}%" if day_ll1>0 else '0%'
    pct2t = f"{round(day_ll2b/day_ll2*100,1)}%" if day_ll2>0 else '0%'
    for ci,val in enumerate(['TOTAL DIA', day_ll1, day_ll1b, pct1t, day_ll2, day_ll2b, pct2t, day_ll1+day_ll2, day_ll1b+day_ll2b],1):
        c = ws1.cell(row=r, column=ci, value=val)
        st(c, bold=True, sz=9, col=WHITE, bg=G_M, ha='center'); c.border = bdr2
    ws1.row_dimensions[r].height = 18; r+=2

# ── HOJA 2: DESGLOSE POR ESTATUS ─────────────────────────────────
ws2 = wb.create_sheet("Estatus por Llamada")
NCOLS2 = 4
trow(ws2, 1, 'WINBIOTA - Desglose por Estatus', NCOLS2)
r2 = 3

section(ws2, r2, 'ESTATUS LLAMADA 1', NCOLS2); r2+=1
hrow(ws2, r2, ['Estatus','Cantidad','% del total','Tipo'], [30,12,14,12]); r2+=1
total_e1 = sum(estatus1_counts.values())
for estatus, cnt in sorted(estatus1_counts.items(), key=lambda x: -x[1]):
    bg = G_L if r2%2==0 else WHITE
    is_bueno = estatus.upper() in buenos_vals1
    tipo = 'Efectiva' if is_bueno else 'No efectiva'
    bg = 'E8F5E9' if is_bueno else 'FFEBEE'
    pct = f"{round(cnt/total_e1*100,1)}%" if total_e1>0 else '0%'
    for ci,val in enumerate([estatus, cnt, pct, tipo],1):
        c = ws2.cell(row=r2, column=ci, value=val)
        st(c, sz=9, bg=bg, ha='center' if ci>1 else 'left'); c.border = bdr2
    ws2.row_dimensions[r2].height = 16; r2+=1
for ci,val in enumerate(['TOTAL', total_e1, '100%', ''],1):
    c = ws2.cell(row=r2, column=ci, value=val)
    st(c, bold=True, sz=9, col=WHITE, bg=G_M, ha='center'); c.border = bdr2
ws2.row_dimensions[r2].height = 18; r2+=2

section(ws2, r2, 'ESTATUS LLAMADA 2', NCOLS2); r2+=1
hrow(ws2, r2, ['Estatus','Cantidad','% del total','Tipo'], [30,12,14,12]); r2+=1
total_e2 = sum(estatus2_counts.values())
for estatus, cnt in sorted(estatus2_counts.items(), key=lambda x: -x[1]):
    is_bueno = estatus.upper() in buenos_vals2
    tipo = 'Efectiva' if is_bueno else 'No efectiva'
    bg = 'E8F5E9' if is_bueno else 'FFEBEE'
    pct = f"{round(cnt/total_e2*100,1)}%" if total_e2>0 else '0%'
    for ci,val in enumerate([estatus, cnt, pct, tipo],1):
        c = ws2.cell(row=r2, column=ci, value=val)
        st(c, sz=9, bg=bg, ha='center' if ci>1 else 'left'); c.border = bdr2
    ws2.row_dimensions[r2].height = 16; r2+=1
for ci,val in enumerate(['TOTAL', total_e2, '100%', ''],1):
    c = ws2.cell(row=r2, column=ci, value=val)
    st(c, bold=True, sz=9, col=WHITE, bg=G_M, ha='center'); c.border = bdr2
ws2.row_dimensions[r2].height = 18

# ── HOJA 3: LEADS EFECTIVOS ───────────────────────────────────────
ws3 = wb.create_sheet("Leads Efectivos")
NCOLS3 = 5
trow(ws3, 1, 'WINBIOTA - Leads Efectivos (buenos estatus)', NCOLS3)
hrow(ws3, 2, ['Fecha Lead','Nombre','Telefono','Estatus','Llamada'], [14,30,16,24,10])
for ri, row in enumerate(sorted(leads_efectivos, key=lambda x: x[0]), 3):
    bg = G_L if ri%2==0 else WHITE
    for ci, val in enumerate(row, 1):
        c = ws3.cell(row=ri, column=ci, value=val)
        st(c, sz=9, bg=bg); c.border = bdr2
    ws3.row_dimensions[ri].height = 16
ws3.freeze_panes = 'A3'

wb.save(OUTPUT)
print(f"Excel guardado. Dias={len(by_date)}, Leads efectivos={len(leads_efectivos)}")

# ── Send email ────────────────────────────────────────────────────
now = datetime.datetime.utcnow()
now_es = (now + datetime.timedelta(hours=1)).strftime('%d/%m/%Y %H:%M')
msg = MIMEMultipart()
msg['From'] = EMAIL_USER
msg['To']   = EMAIL_USER
msg['Subject'] = f'Winbiota - Reporte Diario {now_es}'
msg.attach(MIMEText(
    f"Reporte diario generado el {now_es}.\n"
    f"Dias con datos: {len(by_date)}\n"
    f"Leads efectivos historico: {len(leads_efectivos)}", 'plain'))
with open(OUTPUT, 'rb') as f:
    part = MIMEBase('application', 'octet-stream')
    part.set_payload(f.read())
encoders.encode_base64(part)
part.add_header('Content-Disposition',
    f'attachment; filename="Winbiota_Diario_{now.strftime("%Y%m%d")}.xlsx"')
msg.attach(part)
with smtplib.SMTP_SSL('smtp.gmail.com', 465) as server:
    server.login(EMAIL_USER, EMAIL_PASS)
    server.sendmail(EMAIL_USER, EMAIL_USER, msg.as_string())
print("Email enviado.")
